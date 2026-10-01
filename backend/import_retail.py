"""Offline, deterministic import of the attributed UCI Online Retail workbook.

Runtime only needs data/retail_monthly.json. openpyxl is an optional importer
requirement; importing this module or using summarize_rows does not need it.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
from typing import Iterable, Mapping, Any

EXPECTED_HEADERS = (
    'InvoiceNo', 'StockCode', 'Description', 'Quantity', 'InvoiceDate',
    'UnitPrice', 'CustomerID', 'Country',
)
EXPECTED_ROWS = 541_909
ZIP_SHA256 = 'f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b'
WORKBOOK_SHA256 = '43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d'
BEFORE_MONTH = '2011-10'
AFTER_MONTH = '2011-11'
DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / 'data' / 'retail_monthly.json'


def canonical_hash(payload: Mapping[str, Any]) -> str:
    """Hash the complete payload, except its own hash, with ASCII JSON encoding."""
    body = {key: value for key, value in payload.items() if key != 'hash'}
    serialized = json.dumps(body, sort_keys=True, separators=(',', ':'), allow_nan=False)
    return hashlib.sha256(serialized.encode('utf-8')).hexdigest()


def _money(value: Decimal) -> str:
    text = format(value, 'f')
    return text.rstrip('0').rstrip('.') if '.' in text else text


def _number(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return number if number.is_finite() else None


def _missing(value: Any) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        return True
    return isinstance(value, (float, Decimal)) and _number(value) is None


def _empty_reference() -> dict[str, Any]:
    return {'gross_sales_gbp': Decimal(0), 'invoices': set(), 'lines': 0, 'units': 0}


def _reference_view(value: Mapping[str, Any]) -> dict[str, Any]:
    return {
        'gross_sales_gbp': _money(value['gross_sales_gbp']),
        'orders': len(value['invoices']), 'lines': value['lines'], 'units': value['units'],
    }


def summarize_rows(rows: Iterable[Mapping[str, Any]]) -> dict[str, Any]:
    """Aggregate source rows, preserving repetitions and auditing exclusions.

    Only positive non-cancelled lines contribute to gross sales. Missing
    customer IDs/descriptions are counted but do not exclude a sales line.
    Exclusion-reason counts overlap; excluded_lines counts each row once.
    Reference totals fold accepted source rows directly, not the aggregates.
    """
    audit: dict[str, Any] = dict.fromkeys((
        'total_rows', 'cancellation_lines', 'nonpositive_quantity_lines',
        'nonpositive_price_lines', 'missing_customer_id_lines',
        'missing_description_lines', 'exact_duplicate_rows', 'invalid_lines',
        'included_lines', 'excluded_lines', 'country_count',
    ), 0)
    source_dates: list[datetime] = []
    seen: set[tuple[Any, ...]] = set()
    countries: set[str] = set()
    monthly: dict[tuple[str, str], dict[str, Any]] = {}
    reference_totals = {BEFORE_MONTH: _empty_reference(), AFTER_MONTH: _empty_reference()}
    country_reference: dict[str, dict[str, dict[str, Any]]] = defaultdict(
        lambda: {BEFORE_MONTH: _empty_reference(), AFTER_MONTH: _empty_reference()}
    )
    for row in rows:
        audit['total_rows'] += 1
        # Compare all eight original field values, before filtering. Repetitions
        # after the first occurrence are counted, never silently removed.
        raw_key = tuple(row.get(column) for column in EXPECTED_HEADERS)
        if raw_key in seen:
            audit['exact_duplicate_rows'] += 1
        else:
            seen.add(raw_key)
        audit['missing_customer_id_lines'] += int(_missing(row.get('CustomerID')))
        audit['missing_description_lines'] += int(_missing(row.get('Description')))
        raw_invoice = row.get('InvoiceNo')
        invoice = '' if _missing(raw_invoice) else str(raw_invoice).strip()
        cancelled = invoice.upper().startswith('C')
        quantity = _number(row.get('Quantity'))
        price = _number(row.get('UnitPrice'))
        nonpositive_quantity = quantity is not None and quantity <= 0
        nonpositive_price = price is not None and price <= 0
        audit['cancellation_lines'] += int(cancelled)
        audit['nonpositive_quantity_lines'] += int(nonpositive_quantity)
        audit['nonpositive_price_lines'] += int(nonpositive_price)
        raw_date = row.get('InvoiceDate')
        timestamp = raw_date if isinstance(raw_date, datetime) else (
            datetime.combine(raw_date, time()) if isinstance(raw_date, date) else None
        )
        if timestamp is not None and timestamp.tzinfo is not None:
            timestamp = None  # Workbook dates are local, with no timezone supplied.
        if timestamp is not None:
            source_dates.append(timestamp)
        raw_country = row.get('Country')
        country = raw_country.strip() if isinstance(raw_country, str) else ''
        invalid = (
            not invoice or isinstance(raw_invoice, bool) or not country or timestamp is None
            or quantity is None or price is None
            or (quantity is not None and quantity != quantity.to_integral_value())
        )
        if invalid:
            audit['invalid_lines'] += 1
        if invalid or cancelled or nonpositive_quantity or nonpositive_price:
            audit['excluded_lines'] += 1
            continue
        audit['included_lines'] += 1
        countries.add(country)
        units = int(quantity)
        gross = quantity * price
        month = timestamp.strftime('%Y-%m')
        key = (month, country)
        if key not in monthly:
            monthly[key] = {'lines': 0, 'invoices': set(), 'units': 0, 'gross': Decimal(0)}
        group = monthly[key]
        group['lines'] += 1
        group['invoices'].add(invoice)
        group['units'] += units
        group['gross'] += gross
        if month in reference_totals:
            # Separate source-row accumulators: no SQL and no monthly roll-up.
            for accumulator in (reference_totals[month], country_reference[country][month]):
                accumulator['gross_sales_gbp'] += gross
                accumulator['invoices'].add(invoice)
                accumulator['lines'] += 1
                accumulator['units'] += units
    audit['country_count'] = len(countries)
    audit['source_date_min'] = min(source_dates).isoformat() if source_dates else None
    audit['source_date_max'] = max(source_dates).isoformat() if source_dates else None
    before, after = reference_totals[BEFORE_MONTH], reference_totals[AFTER_MONTH]
    payload = {
        'source': {
            'name': 'UCI Online Retail',
            'url': 'https://archive.ics.uci.edu/dataset/352/online+retail',
            'download_url': 'https://archive.ics.uci.edu/static/public/352/online+retail.zip',
            'doi': '10.24432/C5BW33',
            'citation': 'Chen, D. (2015). Online Retail [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5BW33.',
            'license': 'CC BY 4.0',
            'license_url': 'https://creativecommons.org/licenses/by/4.0/',
            'source_sha256': ZIP_SHA256,
            'workbook_sha256': WORKBOOK_SHA256,
            'source_rows': audit['total_rows'],
        },
        'audit': audit,
        'monthly': [
            {'month': month, 'country': country, 'lines': value['lines'],
             'orders': len(value['invoices']), 'units': value['units'],
             'gross_sales_gbp': _money(value['gross'])}
            for (month, country), value in sorted(monthly.items())
        ],
        'reference': {
            'before_month': BEFORE_MONTH, 'after_month': AFTER_MONTH,
            'before': _reference_view(before), 'after': _reference_view(after),
            'delta_gbp': _money(after['gross_sales_gbp'] - before['gross_sales_gbp']),
            'country_totals': [
                {'country': country,
                 'before': _reference_view(values[BEFORE_MONTH]),
                 'after': _reference_view(values[AFTER_MONTH]),
                 'delta_gbp': _money(values[AFTER_MONTH]['gross_sales_gbp'] - values[BEFORE_MONTH]['gross_sales_gbp'])}
                for country, values in sorted(country_reference.items())
            ],
        },
    }
    payload['hash'] = canonical_hash(payload)
    return payload



def validate_public_aggregates(payload: Mapping[str, Any]) -> None:
    """Allow only non-identifying aggregate fields in every output data row."""
    measures = {'gross_sales_gbp', 'orders', 'lines', 'units'}
    if any(set(row) != measures | {'month', 'country'} for row in payload['monthly']):
        raise ValueError('Unexpected public aggregate fields; refusing to save')
    reference = payload['reference']
    sides = [reference['before'], reference['after']]
    for row in reference['country_totals']:
        if set(row) != {'country', 'before', 'after', 'delta_gbp'}:
            raise ValueError('Unexpected public aggregate fields; refusing to save')
        sides.extend((row['before'], row['after']))
    if any(set(side) != measures for side in sides):
        raise ValueError('Unexpected public aggregate fields; refusing to save')


def validate_workbook_contract(headers: Iterable[str], max_row: int, max_column: int) -> None:
    if tuple(headers) != EXPECTED_HEADERS:
        raise ValueError('Workbook headers do not match the eight UCI Online Retail columns')
    if (max_row, max_column) != (EXPECTED_ROWS + 1, len(EXPECTED_HEADERS)):
        raise ValueError('Workbook dimensions must be 541910 rows (including header) by 8 columns')


def import_workbook(path: Path) -> dict[str, Any]:
    """Read every source row and reject a workbook that is not the pinned source."""
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != WORKBOOK_SHA256:
        raise ValueError(f'Workbook SHA-256 does not match the pinned UCI source: {digest}')
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit('Install the optional importer dependency: pip install -r requirements-data.txt') from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ['Online Retail']:
            raise ValueError('Expected one sheet named Online Retail')
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = tuple(next(values))
        validate_workbook_contract(headers, sheet.max_row, sheet.max_column)
        payload = summarize_rows(dict(zip(headers, values_row, strict=True)) for values_row in values)
        if payload['audit']['total_rows'] != EXPECTED_ROWS:
            raise ValueError('Parsed source row count does not match the workbook contract')
        validate_public_aggregates(payload)
        return payload
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('xlsx', type=Path, help='Downloaded official Online Retail.xlsx; never fetched at runtime')
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    payload = import_workbook(args.xlsx)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, sort_keys=True, separators=(',', ':'), allow_nan=False) + '\n', encoding='utf-8')
    print(json.dumps({'output': str(args.output), 'hash': payload['hash'], 'audit': payload['audit'],
                      'monthly_groups': len(payload['monthly']), 'before': payload['reference']['before'],
                      'after': payload['reference']['after'], 'delta_gbp': payload['reference']['delta_gbp']}, indent=2))


if __name__ == '__main__':
    main()
